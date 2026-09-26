import csv
import hashlib
import io
import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from psycopg2.extras import Json

from db import get_db_connection
from services.client_service import find_possible_client_duplicates
from services.client_information_labels import (
    get_client_information_label_definitions,
    get_client_information_labels,
    get_client_record_note_label_definitions,
    get_client_record_note_labels,
)


from services.peachpos_processor_labels import (
    get_peachpos_processor_field_definitions,
    get_peachpos_processor_field_labels,
)


class ImportServiceError(Exception):
    """Raised when a PSP import file cannot be processed safely."""


IMPORT_MAX_BYTES = 10 * 1024 * 1024
IMPORT_MAX_ROWS = 25000
IMPORT_MAX_COLUMNS = 200

IMPORT_ALLOWED_EXTENSIONS = {
    "csv",
    "xlsx",
}


def _clean_cell_value(value):
    """
    Convert imported spreadsheet values into safe, predictable
    values for preview/mapping without changing their meaning.
    """

    if value is None:
        return ""

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    return str(value).strip()


def _trim_trailing_empty(values):
    values = list(values)

    while values and not str(values[-1] or "").strip():
        values.pop()

    return values


def validate_import_upload(uploaded_file):
    """
    Validate a customer-supplied CSV/XLSX upload before parsing.

    Returns the lowercase extension when valid.
    Raises ImportServiceError when invalid.
    """

    if uploaded_file is None:
        raise ImportServiceError(
            "Choose a CSV or Excel file to import."
        )

    filename = str(
        uploaded_file.filename or ""
    ).strip()

    if not filename:
        raise ImportServiceError(
            "Choose a CSV or Excel file to import."
        )

    if "." not in filename:
        raise ImportServiceError(
            "The import filename must include a valid extension."
        )

    extension = (
        filename.rsplit(".", 1)[1]
        .strip()
        .lower()
    )

    if extension not in IMPORT_ALLOWED_EXTENSIONS:
        raise ImportServiceError(
            "Import files must be CSV (.csv) or Excel (.xlsx)."
        )

    uploaded_file.stream.seek(
        0,
        io.SEEK_END,
    )

    file_size = uploaded_file.stream.tell()

    uploaded_file.stream.seek(0)

    if file_size <= 0:
        raise ImportServiceError(
            "The selected import file is empty."
        )

    if file_size > IMPORT_MAX_BYTES:
        raise ImportServiceError(
            "Import files must be 10 MB or smaller."
        )

    return extension


def _parse_csv(payload):
    """
    Parse CSV bytes using common export encodings.
    """

    decoded = None

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
    ):
        try:
            decoded = payload.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        raise ImportServiceError(
            "The CSV file encoding could not be read."
        )

    stream = io.StringIO(decoded)

    try:
        rows = list(csv.reader(stream))
    except csv.Error as exc:
        raise ImportServiceError(
            "The CSV file could not be parsed."
        ) from exc

    return rows


def _parse_xlsx(payload):
    """
    Parse the active worksheet from an XLSX workbook.
    """

    try:
        workbook = load_workbook(
            filename=io.BytesIO(payload),
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise ImportServiceError(
            "The Excel workbook could not be read."
        ) from exc

    try:
        worksheet = workbook.active

        rows = [
            list(row)
            for row in worksheet.iter_rows(
                values_only=True
            )
        ]
    finally:
        workbook.close()

    return rows


def parse_import_upload(uploaded_file):
    """
    Parse one CSV/XLSX upload into a generic import-table structure.

    No database writes occur here. The returned rows are intentionally
    entity-neutral so Clients, Income, Expenses, Loans, Appointments,
    and future import profiles can all use the same engine.
    """

    extension = validate_import_upload(
        uploaded_file
    )

    filename = Path(
        str(uploaded_file.filename or "")
    ).name

    uploaded_file.stream.seek(0)
    payload = uploaded_file.stream.read()
    uploaded_file.stream.seek(0)

    if extension == "csv":
        raw_rows = _parse_csv(payload)
    else:
        raw_rows = _parse_xlsx(payload)

    normalized_rows = []

    for raw_row in raw_rows:
        values = _trim_trailing_empty(
            [
                _clean_cell_value(value)
                for value in raw_row
            ]
        )

        if not values:
            continue

        normalized_rows.append(values)

    if not normalized_rows:
        raise ImportServiceError(
            "The import file does not contain any usable rows."
        )

    headers = normalized_rows[0]

    if not any(
        str(header or "").strip()
        for header in headers
    ):
        raise ImportServiceError(
            "The import file must contain a header row."
        )

    if len(headers) > IMPORT_MAX_COLUMNS:
        raise ImportServiceError(
            f"Import files may contain no more than "
            f"{IMPORT_MAX_COLUMNS} columns."
        )

    headers = [
        (
            str(header).strip()
            if str(header or "").strip()
            else f"Column {index + 1}"
        )
        for index, header in enumerate(headers)
    ]

    source_rows = normalized_rows[1:]

    if len(source_rows) > IMPORT_MAX_ROWS:
        raise ImportServiceError(
            f"Import files may contain no more than "
            f"{IMPORT_MAX_ROWS:,} data rows."
        )

    rows = []

    for source_index, source_row in enumerate(
        source_rows,
        start=2,
    ):
        values = list(source_row)

        if len(values) < len(headers):
            values.extend(
                [""] * (
                    len(headers) - len(values)
                )
            )
        elif len(values) > len(headers):
            raise ImportServiceError(
                f"Row {source_index} contains more columns "
                "than the header row."
            )

        rows.append({
            "row_number": source_index,
            "values": values,
        })

    return {
        "filename": filename,
        "extension": extension,
        "headers": headers,
        "rows": rows,
        "row_count": len(rows),
        "column_count": len(headers),
    }


# =========================================================
# CLIENT IMPORT PROFILE
# =========================================================


CLIENT_IMPORT_COMMUNICATION_DEFAULTS = {
    "ok_to_call": False,
    "ok_to_text": False,
    "ok_to_email": False,
    "sms_opt_in": False,
    "sms_opt_out": False,
    "email_opt_in": False,
    "email_opt_out": False,
    "sms_marketing_status": "opt_out",
    "email_marketing_status": "not_set",
    "sms_consent_timestamp": None,
    "email_consent_timestamp": None,
}


CLIENT_INFORMATION_IMPORT_FIELDS = tuple(
    {
        "key": definition["field_key"],
        "label": definition["default_label"],
        "required": False,
        "aliases": (),
        "field_type": definition["label_type"],
    }
    for definition in get_client_information_label_definitions()
    if definition.get("label_type") != "section"
)


CLIENT_IMPORT_FIELDS = (
    {
        "key": "first_name",
        "label": "First Name",
        "required": True,
        "aliases": (
            "first",
            "firstname",
            "first name",
            "given name",
            "given_name",
            "fname",
        ),
    },
    {
        "key": "last_name",
        "label": "Last Name",
        "required": True,
        "aliases": (
            "last",
            "lastname",
            "last name",
            "family name",
            "family_name",
            "surname",
            "lname",
        ),
    },
    {
        "key": "phone",
        "label": "Phone",
        "required": False,
        "aliases": (
            "phone",
            "phone number",
            "telephone",
            "mobile",
            "mobile phone",
            "cell",
            "cell phone",
        ),
    },
    {
        "key": "email",
        "label": "Email",
        "required": False,
        "aliases": (
            "email",
            "email address",
            "e-mail",
            "e-mail address",
        ),
    },
    {
        "key": "birth_date",
        "label": "Birth Date",
        "required": False,
        "aliases": (
            "birth date",
            "birthdate",
            "birthday",
            "date of birth",
            "dob",
        ),
    },
    {
        "key": "address",
        "label": "Address",
        "required": False,
        "aliases": (
            "address",
            "street",
            "street address",
            "address 1",
            "address1",
        ),
    },
    {
        "key": "city",
        "label": "City",
        "required": False,
        "aliases": (
            "city",
            "town",
        ),
    },
    {
        "key": "state",
        "label": "State",
        "required": False,
        "aliases": (
            "state",
            "province",
            "region",
        ),
    },
    {
        "key": "zip",
        "label": "ZIP / Postal Code",
        "required": False,
        "aliases": (
            "zip",
            "zip code",
            "zipcode",
            "postal code",
            "postalcode",
        ),
    },
    {
        "key": "emergency_contact_name",
        "label": "Emergency Contact Name",
        "required": False,
        "aliases": (
            "emergency contact",
            "emergency contact name",
            "emergency name",
        ),
    },
    {
        "key": "emergency_contact_phone",
        "label": "Emergency Contact Phone",
        "required": False,
        "aliases": (
            "emergency contact phone",
            "emergency phone",
        ),
    },
    {
        "key": "client_status",
        "label": "Client Status",
        "required": False,
        "aliases": (
            "client status",
            "customer status",
            "status",
        ),
    },
    {
        "key": "preferred_language",
        "label": "Preferred Language",
        "required": False,
        "aliases": (
            "preferred language",
            "language",
        ),
    },
    {
        "key": "preferred_contact_method",
        "label": "Preferred Contact Method",
        "required": False,
        "aliases": (
            "preferred contact method",
            "contact method",
            "preferred contact",
        ),
    },
    {
        "key": "notes_one",
        "label": "Notes One",
        "required": False,
        "aliases": (
            "notes",
            "note",
            "client notes",
            "customer notes",
            "notes one",
            "client note 1",
            "client note one",
        ),
    },
    {
        "key": "notes_two",
        "label": "Notes Two",
        "required": False,
        "aliases": (
            "notes two",
            "client note 2",
            "client note two",
            "secondary notes",
        ),
    },
    {
        "key": "notes_three",
        "label": "Notes Three",
        "required": False,
        "aliases": (
            "notes three",
            "client note 3",
            "client note three",
        ),
    },
    {
        "key": "active_client",
        "label": "Active Client",
        "required": False,
        "aliases": (
            "active",
            "active client",
            "active customer",
            "is active",
        ),
    },
    *CLIENT_INFORMATION_IMPORT_FIELDS,
)


# =========================================================
# PEACHPOS INCOME IMPORT PROFILE
# =========================================================


PEACHPOS_PROCESSOR_IMPORT_FIELDS = tuple(
    {
        "key": definition["field_key"],
        "label": definition["default_label"],
        "required": False,
        "aliases": (),
        "field_type": definition["label_type"],
    }
    for definition in get_peachpos_processor_field_definitions()
)


PEACHPOS_INCOME_IMPORT_FIELDS = (
    {
        "key": "transaction_at",
        "label": "Transaction Date / Time",
        "required": True,
        "aliases": (
            "transaction date",
            "transaction datetime",
            "transaction date time",
            "date time",
            "datetime",
            "sale date",
            "created at",
            "created_at",
        ),
    },
    {
        "key": "transaction_time",
        "label": "Transaction Time",
        "required": False,
        "aliases": (
            "transaction time",
            "time",
            "sale time",
            "payment time",
            "created time",
        ),
    },
    {
        "key": "processor_payment_id",
        "label": "Processor Transaction ID",
        "required": True,
        "aliases": (
            "transaction id",
            "transaction_id",
            "payment id",
            "payment_id",
            "processor transaction id",
            "processor payment id",
            "reference id",
            "reference number",
            "transaction number",
        ),
    },
    {
        "key": "merchant_account_identifier",
        "label": "Merchant ID / Account ID",
        "required": True,
        "aliases": (
            "merchant id",
            "merchant_id",
            "merchant account id",
            "merchant account",
            "merchant account identifier",
            "mid",
            "account id",
            "account_id",
            "processor account id",
        ),
    },
    {
        "key": "pos_amount",
        "label": "Sales",
        "required": True,
        "aliases": (
            "sales",
            "sale amount",
            "sales amount",
            "subtotal",
            "pre tax sales",
            "pre-tax sales",
            "net sales before tax",
        ),
    },
    {
        "key": "tax_amount",
        "label": "Tax",
        "required": False,
        "aliases": (
            "tax",
            "tax amount",
            "sales tax",
            "tax collected",
        ),
    },
    {
        "key": "tip_amount",
        "label": "Tips",
        "required": False,
        "aliases": (
            "tip",
            "tips",
            "tip amount",
            "tips collected",
            "gratuity",
        ),
    },
    {
        "key": "total_amount",
        "label": "Gross Collected",
        "required": False,
        "aliases": (
            "total",
            "total amount",
            "gross",
            "gross amount",
            "gross collected",
            "amount",
            "payment amount",
        ),
    },
    {
        "key": "processing_fee_amount",
        "label": "Processor Fees",
        "required": False,
        "aliases": (
            "fee",
            "fees",
            "processing fee",
            "processing fees",
            "processor fee",
            "processor fees",
        ),
    },
    {
        "key": "net_received",
        "label": "Net Received",
        "required": False,
        "aliases": (
            "net",
            "net amount",
            "net received",
            "deposit amount",
            "settlement amount",
        ),
    },
    {
        "key": "payment_method",
        "label": "Payment Method",
        "required": False,
        "aliases": (
            "payment method",
            "payment type",
            "tender",
            "tender type",
            "card type",
        ),
    },
    {
        "key": "processor_status",
        "label": "Processor Status",
        "required": False,
        "aliases": (
            "status",
            "payment status",
            "transaction status",
            "processor status",
        ),
    },
    *PEACHPOS_PROCESSOR_IMPORT_FIELDS,
)


def normalize_import_header(value):
    """
    Normalize a source column heading only for matching purposes.
    The original heading remains unchanged for display.
    """

    value = str(value or "").strip().lower()

    return "".join(
        character
        for character in value
        if character.isalnum()
    )


def get_import_profile(entity_type):
    """
    Return the entity-specific rules used by the shared Import Engine.
    """

    entity_type = str(
        entity_type or ""
    ).strip().lower()

    if entity_type == "clients":
        return {
            "entity_type": "clients",
            "display_name": "Clients",
            "fields": CLIENT_IMPORT_FIELDS,
            "defaults": {
                "client_status": "Current",
                "active_client": True,
            },
        }

    if entity_type == "peachpos_income":
        return {
            "entity_type": "peachpos_income",
            "display_name": "PeachPOS Income",
            "fields": PEACHPOS_INCOME_IMPORT_FIELDS,
            "defaults": {},
        }

    raise ImportServiceError(
        f"Unsupported import type: {entity_type or 'unknown'}."
    )



def get_workspace_import_profile(
    cur,
    *,
    spa_id,
    business_unit_id,
    entity_type,
    options=None,
):
    """
    Return a workspace-aware copy of an Import Engine profile.

    Backend target keys remain immutable. Workspace customization
    affects only display labels and exact-header matching aliases.
    """

    base_profile = get_import_profile(
        entity_type
    )

    profile = {
        **base_profile,
        "fields": [
            dict(field)
            for field in base_profile["fields"]
        ],
    }

    if profile["entity_type"] == "peachpos_income":
        options_payload = dict(options or {})
        credit_processor_id = options_payload.get(
            "credit_processor_id"
        )

        try:
            credit_processor_id = int(
                credit_processor_id
            )
        except (TypeError, ValueError):
            raise ImportServiceError(
                "Choose a credit processor before importing "
                "PeachPOS Income."
            )

        if credit_processor_id <= 0:
            raise ImportServiceError(
                "Choose a credit processor before importing "
                "PeachPOS Income."
            )

        try:
            resolved_processor_labels = (
                get_peachpos_processor_field_labels(
                    cur,
                    spa_id=spa_id,
                    business_unit_id=business_unit_id,
                    credit_processor_id=credit_processor_id,
                    require_active=True,
                )
            )
        except ValueError as exc:
            raise ImportServiceError(str(exc)) from exc

        processor_definitions = {
            definition["field_key"]: definition
            for definition
            in get_peachpos_processor_field_definitions()
        }

        for field in profile["fields"]:
            field_key = field["key"]
            definition = processor_definitions.get(
                field_key
            )
            if not definition:
                continue

            current_label = (
                resolved_processor_labels.get(
                    field_key
                )
                or definition["default_label"]
            )
            field["label"] = current_label
            field["workspace_match_labels"] = (
                current_label,
                definition["default_label"],
            )

        return profile

    if profile["entity_type"] != "clients":
        return profile

    resolved_client_information_labels = (
        get_client_information_labels(
            cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
        )
    )

    client_information_definitions = {
        definition["field_key"]: definition
        for definition
        in get_client_information_label_definitions()
        if definition.get("label_type") != "section"
    }

    resolved_note_labels = (
        get_client_record_note_labels(
            cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
        )
    )

    note_definitions = {
        definition["field_key"]: definition
        for definition
        in get_client_record_note_label_definitions()
    }

    workspace_definitions = {
        **client_information_definitions,
        **note_definitions,
    }

    resolved_workspace_labels = {
        **resolved_client_information_labels,
        **resolved_note_labels,
    }

    for field in profile["fields"]:
        field_key = field["key"]

        definition = workspace_definitions.get(
            field_key
        )

        if not definition:
            continue

        current_label = (
            resolved_workspace_labels.get(
                field_key
            )
            or definition["default_label"]
        )

        field["label"] = current_label

        field["workspace_match_labels"] = (
            current_label,
            definition["default_label"],
            definition["legacy_label"],
        )

    return profile

def suggest_import_mapping(
    headers,
    entity_type="clients",
    *,
    profile=None,
):
    """
    Suggest safe source-column -> PSP-field mappings.

    Matching is exact after header normalization.

    Priority:
      1. current workspace label
      2. generic PSP label
      3. original / legacy PSP label
      4. immutable backend key and known aliases

    A priority level must resolve to exactly one target.
    Ambiguous or repeated matches remain unmapped for review.
    """

    if profile is None:
        profile = get_import_profile(
            entity_type
        )

    match_levels = [
        {},
        {},
        {},
        {},
    ]

    for field in profile["fields"]:
        field_key = field["key"]

        workspace_labels = tuple(
            field.get(
                "workspace_match_labels",
                (),
            )
        )

        current_label = (
            workspace_labels[0]
            if len(workspace_labels) > 0
            else field["label"]
        )

        generic_label = (
            workspace_labels[1]
            if len(workspace_labels) > 1
            else field["label"]
        )

        legacy_label = (
            workspace_labels[2]
            if len(workspace_labels) > 2
            else None
        )

        level_values = (
            (current_label,),
            (generic_label,),
            (legacy_label,),
            (
                field_key,
                *field.get("aliases", ()),
            ),
        )

        for level_index, aliases in enumerate(
            level_values
        ):
            for alias in aliases:
                normalized = normalize_import_header(
                    alias
                )

                if not normalized:
                    continue

                match_levels[level_index].setdefault(
                    normalized,
                    set(),
                ).add(
                    field_key
                )

    suggestions = []
    used_targets = set()

    for source_index, header in enumerate(headers):
        normalized_header = normalize_import_header(
            header
        )

        target_key = None

        for lookup in match_levels:
            candidates = lookup.get(
                normalized_header,
                set(),
            )

            if not candidates:
                continue

            if len(candidates) == 1:
                candidate = next(
                    iter(candidates)
                )

                if candidate not in used_targets:
                    target_key = candidate
                    used_targets.add(candidate)

            # A match at a higher-priority level prevents
            # falling through to lower-priority aliases.
            break

        suggestions.append({
            "source_index": source_index,
            "source_header": str(
                header or ""
            ).strip(),
            "target_key": target_key,
        })

    return suggestions


# =========================================================
# CLIENT IMPORT ROW NORMALIZATION / VALIDATION
# =========================================================

def _normalize_import_email(value):
    return str(value or "").strip().lower()


def _is_valid_import_email(value):
    value = _normalize_import_email(value)

    if not value:
        return True

    if value.count("@") != 1:
        return False

    local_part, domain = value.split("@", 1)

    return bool(
        local_part
        and domain
        and "." in domain
        and not any(
            character.isspace()
            for character in value
        )
    )


def _parse_import_money(value, *, default=None):
    """
    Normalize common processor-export currency values.
    Returns Decimal for nonblank values.
    Blank values return the supplied default.
    """
    if value is None:
        return default

    if isinstance(value, Decimal):
        return value

    raw = str(value).strip()
    if not raw:
        return default

    negative_parentheses = (
        raw.startswith("(")
        and raw.endswith(")")
    )

    if negative_parentheses:
        raw = raw[1:-1].strip()

    raw = (
        raw.replace(",", "")
        .replace("$", "")
        .strip()
    )

    try:
        amount = Decimal(raw)
    except (InvalidOperation, ValueError):
        raise ImportServiceError(
            f"Amount '{value}' is not a valid monetary value."
        )

    if negative_parentheses:
        amount = -abs(amount)

    return amount.quantize(Decimal("0.01"))


def _parse_import_transaction_datetime(
    value,
    *,
    time_value=None,
):
    """
    Normalize a processor transaction date/time.

    Supports either one combined timestamp or a separate
    date column plus time column. Naive timestamps remain
    naive here; posting later applies the workspace timezone.
    PSP does not invent a missing transaction time.
    """
    raw = str(value or "").strip()
    raw_time = str(time_value or "").strip()

    if not raw:
        raise ImportServiceError(
            "Transaction Date / Time is required."
        )

    if raw_time:
        raw = f"{raw} {raw_time}"

    iso_candidate = raw
    if iso_candidate.endswith("Z"):
        iso_candidate = iso_candidate[:-1] + "+00:00"

    try:
        parsed = datetime.fromisoformat(iso_candidate)
        if isinstance(parsed, datetime):
            return parsed.isoformat()
    except ValueError:
        pass

    datetime_formats = (
        "%m/%d/%Y %I:%M:%S %p",
        "%m/%d/%Y %I:%M %p",
        "%m/%d/%y %I:%M:%S %p",
        "%m/%d/%y %I:%M %p",
        "%m-%d-%Y %I:%M:%S %p",
        "%m-%d-%Y %I:%M %p",
        "%m/%d/%Y %H:%M:%S",
        "%m/%d/%Y %H:%M",
        "%m/%d/%y %H:%M:%S",
        "%m/%d/%y %H:%M",
        "%m-%d-%Y %H:%M:%S",
        "%m-%d-%Y %H:%M",
        "%Y-%m-%d %I:%M:%S %p",
        "%Y-%m-%d %I:%M %p",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
    )

    for date_format in datetime_formats:
        try:
            parsed = datetime.strptime(
                raw,
                date_format,
            )
            return parsed.isoformat()
        except ValueError:
            continue

    date_only_formats = (
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%m-%d-%Y",
        "%m-%d-%y",
    )

    for date_format in date_only_formats:
        try:
            datetime.strptime(raw, date_format)
        except ValueError:
            continue
        raise ImportServiceError(
            "Transaction time is missing. Map the processor "
            "time column or use a combined Date / Time column."
        )

    raise ImportServiceError(
        f"Transaction Date / Time '{raw}' is not in a "
        "supported format."
    )


def _parse_import_date(value):
    """
    Accept common customer-export date formats and return ISO YYYY-MM-DD.
    Blank values remain blank.
    """

    raw = str(value or "").strip()

    if not raw:
        return ""

    date_formats = (
        "%Y-%m-%d",
        "%m/%d/%Y",
        "%m/%d/%y",
        "%m-%d-%Y",
        "%m-%d-%y",
        "%Y/%m/%d",
        "%B %d, %Y",
        "%b %d, %Y",
    )

    for date_format in date_formats:
        try:
            parsed = datetime.strptime(
                raw,
                date_format,
            ).date()

            return parsed.isoformat()
        except ValueError:
            continue

    raise ImportServiceError(
        f"Date '{raw}' is not in a supported format."
    )


def _parse_import_boolean(
    value,
    *,
    default=None,
):
    """
    Parse common boolean values found in CSV/Excel exports.
    """

    if isinstance(value, bool):
        return value

    raw = str(value or "").strip().lower()

    if not raw:
        return default

    truthy = {
        "1",
        "true",
        "yes",
        "y",
        "active",
    }

    falsey = {
        "0",
        "false",
        "no",
        "n",
        "inactive",
    }

    if raw in truthy:
        return True

    if raw in falsey:
        return False

    raise ImportServiceError(
        f"Value '{value}' is not a recognized yes/no value."
    )


def apply_import_mapping(
    values,
    mapping,
    *,
    entity_type="clients",
):
    """
    Convert one source row into an entity-keyed dictionary using
    the reviewed source-column mapping.

    Mapping entries use:
      source_index
      target_key

    Unmapped source columns are intentionally ignored.
    """

    profile = get_import_profile(
        entity_type
    )

    valid_target_keys = {
        field["key"]
        for field in profile["fields"]
    }

    result = dict(
        profile.get(
            "defaults",
            {}
        )
    )

    used_targets = set()

    for item in mapping:
        target_key = str(
            item.get("target_key") or ""
        ).strip()

        if not target_key:
            continue

        if target_key not in valid_target_keys:
            raise ImportServiceError(
                f"Unknown {profile['display_name']} import "
                f"field: {target_key}."
            )

        if target_key in used_targets:
            raise ImportServiceError(
                f"The import mapping assigns more than one "
                f"source column to {target_key}."
            )

        try:
            source_index = int(
                item.get("source_index")
            )
        except (TypeError, ValueError) as exc:
            raise ImportServiceError(
                "The import mapping contains an invalid "
                "source column."
            ) from exc

        if (
            source_index < 0
            or source_index >= len(values)
        ):
            raise ImportServiceError(
                "The import mapping references a source "
                "column outside this row."
            )

        result[target_key] = str(
            values[source_index] or ""
        ).strip()

        used_targets.add(
            target_key
        )

    return result


def validate_client_import_row(
    row_data,
):
    """
    Normalize and validate one mapped Client import row.

    Import policy:
      - first and last name are required
      - phone and email may be blank for historical records
      - email is validated only when present
      - birth date is normalized when present
      - Client Status defaults to Current
      - Active Client defaults to True
      - marketing/SMS/email consent is NOT inferred here
    """

    normalized = {
        key: str(
            value or ""
        ).strip()
        for key, value in row_data.items()
        if key != "active_client"
    }

    errors = []

    first_name = normalized.get(
        "first_name",
        "",
    )

    last_name = normalized.get(
        "last_name",
        "",
    )

    if not first_name:
        errors.append(
            "First Name is required."
        )

    if not last_name:
        errors.append(
            "Last Name is required."
        )

    email = _normalize_import_email(
        normalized.get(
            "email",
            "",
        )
    )

    if email and not _is_valid_import_email(email):
        errors.append(
            "Email address is not valid."
        )

    normalized["email"] = email

    birth_date = normalized.get(
        "birth_date",
        "",
    )

    if birth_date:
        try:
            normalized["birth_date"] = (
                _parse_import_date(
                    birth_date
                )
            )
        except ImportServiceError as exc:
            errors.append(
                str(exc)
            )
            normalized["birth_date"] = (
                birth_date
            )

    user_def_boolean_keys = (
        "recent_injections",
        "recent_laser",
        "pregnant",
        "nursing",
        "using_retinol",
        "using_accutane",
    )

    for field_key in user_def_boolean_keys:
        if field_key not in row_data:
            continue

        try:
            normalized[field_key] = (
                _parse_import_boolean(
                    row_data.get(field_key),
                    default=None,
                )
            )
        except ImportServiceError as exc:
            errors.append(
                str(exc)
            )
            normalized[field_key] = str(
                row_data.get(field_key) or ""
            ).strip()

    if "last_facial_date" in row_data:
        last_facial_date = str(
            row_data.get("last_facial_date") or ""
        ).strip()

        if last_facial_date:
            try:
                normalized["last_facial_date"] = (
                    _parse_import_date(
                        last_facial_date
                    )
                )
            except ImportServiceError as exc:
                errors.append(
                    str(exc)
                )
                normalized["last_facial_date"] = (
                    last_facial_date
                )
        else:
            normalized["last_facial_date"] = ""

    normalized["client_status"] = (
        normalized.get(
            "client_status",
            "",
        )
        or "Current"
    )

    try:
        normalized["active_client"] = (
            _parse_import_boolean(
                row_data.get(
                    "active_client"
                ),
                default=True,
            )
        )
    except ImportServiceError as exc:
        errors.append(
            str(exc)
        )
        normalized["active_client"] = True

    return {
        "valid": not errors,
        "data": normalized,
        "errors": errors,
    }


def validate_peachpos_income_import_row(
    row_data,
):
    """
    Normalize and validate one mapped PeachPOS Income row.

    Data-integrity policy:
      - PSP never invents missing processor data
      - transaction date/time, transaction ID, and Sales
        are required
      - optional monetary values remain blank when absent
      - supplied accounting values are verified when enough
        source values exist to perform the check
      - no missing Total, Fees, Net, Tax, or Tips are derived
    """
    normalized = {
        key: str(value or "").strip()
        for key, value in row_data.items()
    }
    errors = []

    # -------------------------------------------------
    # Processor transaction timestamp
    # -------------------------------------------------
    transaction_value = normalized.get(
        "transaction_at",
        "",
    )
    transaction_time = normalized.get(
        "transaction_time",
        "",
    )

    try:
        normalized["transaction_at"] = (
            _parse_import_transaction_datetime(
                transaction_value,
                time_value=transaction_time,
            )
        )
    except ImportServiceError as exc:
        errors.append(str(exc))
        normalized["transaction_at"] = transaction_value

    # transaction_time is an import-only helper. The original
    # source value remains preserved in source_data.
    normalized.pop("transaction_time", None)

    # -------------------------------------------------
    # Processor transaction identity
    # -------------------------------------------------
    processor_payment_id = normalized.get(
        "processor_payment_id",
        "",
    )
    if not processor_payment_id:
        errors.append(
            "Processor Transaction ID is required."
        )

    merchant_account_identifier = normalized.get(
        "merchant_account_identifier",
        "",
    )
    if not merchant_account_identifier:
        errors.append(
            "Merchant ID / Account ID is required."
        )

    # -------------------------------------------------
    # Monetary values
    # -------------------------------------------------
    money_fields = (
        ("pos_amount", "Sales"),
        ("tax_amount", "Tax"),
        ("tip_amount", "Tips"),
        ("total_amount", "Gross Collected"),
        ("processing_fee_amount", "Processor Fees"),
        ("net_received", "Net Received"),
    )

    parsed_money = {}

    for field_key, field_label in money_fields:
        raw_value = normalized.get(field_key, "")

        if field_key == "pos_amount" and not raw_value:
            errors.append(
                "Sales is required."
            )
            continue

        if not raw_value:
            # Never infer zero or calculate a missing value.
            continue

        try:
            amount = _parse_import_money(raw_value)
        except ImportServiceError as exc:
            errors.append(
                f"{field_label}: {exc}"
            )
            continue

        parsed_money[field_key] = amount
        normalized[field_key] = format(amount, ".2f")

    # PeachPOS represents positive completed sales. Refunds,
    # voids, and other lifecycle cases require separate review.
    pos_amount = parsed_money.get("pos_amount")
    if pos_amount is not None and pos_amount <= Decimal("0.00"):
        errors.append(
            "Sales must be greater than zero for a PeachPOS sale."
        )

    tax_amount = parsed_money.get("tax_amount")
    if tax_amount is not None and tax_amount < Decimal("0.00"):
        errors.append(
            "Tax cannot be negative for a PeachPOS sale."
        )

    tip_amount = parsed_money.get("tip_amount")
    if tip_amount is not None and tip_amount < Decimal("0.00"):
        errors.append(
            "Tips cannot be negative for a PeachPOS sale."
        )

    total_amount = parsed_money.get("total_amount")
    if total_amount is not None and total_amount <= Decimal("0.00"):
        errors.append(
            "Gross Collected must be greater than zero when supplied."
        )

    processing_fee = parsed_money.get("processing_fee_amount")
    if processing_fee is not None and processing_fee < Decimal("0.00"):
        errors.append(
            "Processor Fees cannot be negative in PSP. "
            "Review the processor export fee convention."
        )

    net_received = parsed_money.get("net_received")
    if net_received is not None and net_received < Decimal("0.00"):
        errors.append(
            "Net Received cannot be negative for a PeachPOS sale."
        )

    # -------------------------------------------------
    # Verify supplied accounting relationships.
    # Never fill a missing component.
    # -------------------------------------------------
    if all(
        key in parsed_money
        for key in (
            "pos_amount",
            "tax_amount",
            "tip_amount",
            "total_amount",
        )
    ):
        expected_total = (
            parsed_money["pos_amount"]
            + parsed_money["tax_amount"]
            + parsed_money["tip_amount"]
        ).quantize(Decimal("0.01"))

        if expected_total != parsed_money["total_amount"]:
            errors.append(
                "Sales + Tax + Tips does not equal "
                "Gross Collected."
            )

    if all(
        key in parsed_money
        for key in (
            "total_amount",
            "processing_fee_amount",
            "net_received",
        )
    ):
        expected_net = (
            parsed_money["total_amount"]
            - parsed_money["processing_fee_amount"]
        ).quantize(Decimal("0.01"))

        if expected_net != parsed_money["net_received"]:
            errors.append(
                "Gross Collected - Processor Fees does not "
                "equal Net Received."
            )

    return {
        "valid": not errors,
        "data": normalized,
        "errors": errors,
    }



def verify_peachpos_import_merchant_identity(
    cur,
    *,
    spa_id,
    business_unit_id,
    credit_processor_id,
    prepared_rows,
):
    """
    Hard safety gate for generic non-Square PeachPOS imports.

    The uploaded file must contain exactly one Merchant ID /
    Account ID, and it must exactly match the identifier saved
    for the selected processor account.

    Square uses PSP's native Square integration and is not
    permitted through the generic processor-file importer.
    """
    try:
        credit_processor_id = int(
            credit_processor_id
        )
    except (TypeError, ValueError):
        raise ImportServiceError(
            "Choose a processor before importing "
            "PeachPOS Income."
        )

    cur.execute(
        """
        SELECT
            credit_processor_name,
            merchant_account_identifier
        FROM credit_processors
        WHERE credit_processor_id = %s
          AND spa_id = %s
          AND business_unit_id = %s
        LIMIT 1
        """,
        (
            credit_processor_id,
            spa_id,
            business_unit_id,
        ),
    )
    processor = cur.fetchone()

    if not processor:
        raise ImportServiceError(
            "The selected processor was not found in this "
            "Provider Workspace."
        )

    processor_name = str(
        processor[0] or ""
    ).strip()

    if processor_name.lower() == "square":
        raise ImportServiceError(
            "Square transactions use PSP's native Square "
            "integration and cannot be imported through the "
            "generic PeachPOS Income Import."
        )

    configured_identifier = str(
        processor[1] or ""
    ).strip()

    if not configured_identifier:
        raise ImportServiceError(
            "Merchant ID / Account ID is not configured for "
            f"{processor_name or 'this processor'}. "
            "Enter it in Processor Company Setup before "
            "importing."
        )

    detected_identifiers = set()
    missing_rows = []

    for row in prepared_rows:
        data = row.get("data") or {}
        identifier = str(
            data.get(
                "merchant_account_identifier"
            )
            or ""
        ).strip()

        if not identifier:
            missing_rows.append(
                row.get("row_number")
            )
            continue

        detected_identifiers.add(
            identifier
        )

    if missing_rows:
        raise ImportServiceError(
            "Merchant ID / Account ID is missing from one "
            "or more rows in the uploaded processor file. "
            "Import is blocked."
        )

    if len(detected_identifiers) != 1:
        raise ImportServiceError(
            "The uploaded processor file contains multiple "
            "Merchant ID / Account ID values. "
            "Import is blocked."
        )

    detected_identifier = next(
        iter(detected_identifiers)
    )

    if detected_identifier != configured_identifier:
        raise ImportServiceError(
            "Merchant ID / Account ID in the uploaded file "
            "does not match Processor Company Setup. "
            "Import is blocked."
        )

    return {
        "processor_company": processor_name,
        "detected_identifier": detected_identifier,
        "configured_identifier_at_verification": (
            configured_identifier
        ),
        "verified": True,
    }

def _normalize_import_dropdown_value(value):
    return " ".join(
        str(value or "")
        .strip()
        .lower()
        .split()
    )


def validate_client_import_workspace_dropdowns(
    cur,
    *,
    spa_id,
    business_unit_id,
    prepared_rows,
):
    """
    Resolve database-backed Client Information dropdown values.

    User Def 1 / skin_type_id:
      imported Skin Type display value -> active spa skin_type_id

    User Def 2 / fitzpatrick_id:
      imported Fitzpatrick display value -> active spa fitzpatrick_id

    Values are matched exactly after case/whitespace normalization.
    Blank values remain blank. Unknown values are review errors.
    """

    resolved_labels = get_client_information_labels(
        cur,
        spa_id=spa_id,
        business_unit_id=business_unit_id,
    )

    dropdown_specs = (
        {
            "field_key": "skin_type_id",
            "table": "skin_types",
            "id_column": "skin_type_id",
            "value_column": "skin_type_name",
        },
        {
            "field_key": "fitzpatrick_id",
            "table": "fitzpatrick_types",
            "id_column": "fitzpatrick_id",
            "value_column": "fitzpatrick_level",
        },
    )

    for spec in dropdown_specs:
        cur.execute(
            f"""
                SELECT
                    {spec["id_column"]},
                    {spec["value_column"]}
                FROM {spec["table"]}
                WHERE spa_id = %s
                  AND is_active = TRUE
                ORDER BY {spec["id_column"]}
            """,
            (spa_id,),
        )

        options = cur.fetchall()

        option_lookup = {}

        for option_id, option_value in options:
            normalized_value = (
                _normalize_import_dropdown_value(
                    option_value
                )
            )

            if not normalized_value:
                continue

            option_lookup.setdefault(
                normalized_value,
                [],
            ).append(
                option_id
            )

        field_key = spec["field_key"]

        field_label = (
            resolved_labels.get(field_key)
            or field_key
        )

        for row in prepared_rows:
            data = row.get("data") or {}

            if field_key not in data:
                continue

            raw_value = data.get(field_key)

            normalized_value = (
                _normalize_import_dropdown_value(
                    raw_value
                )
            )

            if not normalized_value:
                data[field_key] = ""
                continue

            matches = option_lookup.get(
                normalized_value,
                [],
            )

            if len(matches) == 1:
                data[field_key] = matches[0]
                continue

            errors = list(
                row.get("errors") or []
            )

            if not options:
                errors.append(
                    f"{field_label} has no active dropdown "
                    "options configured. Leave this field "
                    "unmapped or configure its options before "
                    "importing."
                )
            elif len(matches) > 1:
                errors.append(
                    f"{field_label} value '{raw_value}' matches "
                    "more than one active option."
                )
            else:
                errors.append(
                    f"{field_label} value '{raw_value}' does "
                    "not match an active dropdown option."
                )

            row["errors"] = errors
            row["valid"] = False

    return prepared_rows


def prepare_import_rows(
    parsed_import,
    mapping,
    *,
    entity_type="clients",
):
    """
    Apply a reviewed mapping and entity validation to every source row.

    Still performs no database writes.
    """

    prepared_rows = []

    for source_row in parsed_import.get(
        "rows",
        []
    ):
        mapped = apply_import_mapping(
            source_row.get(
                "values",
                []
            ),
            mapping,
            entity_type=entity_type,
        )

        if entity_type == "clients":
            validation = (
                validate_client_import_row(
                    mapped
                )
            )
        elif entity_type == "peachpos_income":
            validation = (
                validate_peachpos_income_import_row(
                    mapped
                )
            )
        else:
            raise ImportServiceError(
                f"Unsupported import type: {entity_type}."
            )

        prepared_rows.append({
            "row_number": source_row.get(
                "row_number"
            ),
            "valid": validation["valid"],
            "data": validation["data"],
            "errors": validation["errors"],
        })

    return prepared_rows


# =========================================================
# CLIENT IMPORT DUPLICATE DETECTION
# =========================================================

def normalize_client_import_name(value):
    return " ".join(
        str(value or "")
        .strip()
        .lower()
        .split()
    )


def normalize_client_import_phone(value):
    digits = re.sub(
        r"[^0-9]",
        "",
        str(value or ""),
    )

    if (
        len(digits) == 11
        and digits.startswith("1")
    ):
        digits = digits[1:]

    return digits


def annotate_client_import_file_duplicates(
    prepared_rows,
):
    """
    Detect duplicates within one uploaded Client file.

    Strong duplicate:
      - exact normalized email
      - exact normalized phone

    Possible duplicate:
      - exact normalized first + last name only

    Invalid rows are left invalid and are not used as duplicate anchors.
    """

    seen_emails = {}
    seen_phones = {}
    seen_names = {}

    annotated_rows = []

    for row in prepared_rows:
        result = dict(row)

        result["duplicate_type"] = None
        result["duplicate_reasons"] = []
        result["duplicate_row_numbers"] = []

        if not result.get("valid"):
            annotated_rows.append(
                result
            )
            continue

        data = result.get(
            "data",
            {},
        )

        row_number = result.get(
            "row_number"
        )

        email = _normalize_import_email(
            data.get(
                "email",
                "",
            )
        )

        phone = normalize_client_import_phone(
            data.get(
                "phone",
                "",
            )
        )

        first_name = normalize_client_import_name(
            data.get(
                "first_name",
                "",
            )
        )

        last_name = normalize_client_import_name(
            data.get(
                "last_name",
                "",
            )
        )

        name_key = (
            first_name,
            last_name,
        )

        strong_matches = []
        possible_matches = []

        if (
            email
            and email in seen_emails
        ):
            strong_matches.append((
                "Email",
                seen_emails[email],
            ))

        if (
            phone
            and phone in seen_phones
        ):
            strong_matches.append((
                "Phone",
                seen_phones[phone],
            ))

        if (
            first_name
            and last_name
            and name_key in seen_names
        ):
            possible_matches.append((
                "Name",
                seen_names[name_key],
            ))

        if strong_matches:
            result["duplicate_type"] = (
                "strong"
            )

            result["duplicate_reasons"] = sorted({
                reason
                for reason, _row
                in strong_matches
            })

            result["duplicate_row_numbers"] = sorted({
                duplicate_row
                for _reason, duplicate_row
                in strong_matches
            })

        elif possible_matches:
            result["duplicate_type"] = (
                "possible"
            )

            result["duplicate_reasons"] = [
                "Name"
            ]

            result["duplicate_row_numbers"] = sorted({
                duplicate_row
                for _reason, duplicate_row
                in possible_matches
            })

        annotated_rows.append(
            result
        )

        # Only the first occurrence becomes the anchor.
        # This keeps repeated duplicates pointing back to the
        # original source row instead of forming a chain.
        if email:
            seen_emails.setdefault(
                email,
                row_number,
            )

        if phone:
            seen_phones.setdefault(
                phone,
                row_number,
            )

        if first_name and last_name:
            seen_names.setdefault(
                name_key,
                row_number,
            )

    return annotated_rows


def annotate_client_import_existing_duplicates(
    cur,
    *,
    spa_id,
    business_unit_id,
    prepared_rows,
):
    """
    Compare prepared Client import rows against existing PSP clients.

    Uses the same authoritative duplicate checker as normal Add/Edit
    Client workflows.

    Strong duplicate:
      - exact normalized email
      - exact normalized phone

    Possible duplicate:
      - exact normalized first + last name only

    This function performs database reads only.
    """

    annotated_rows = []

    for row in prepared_rows:
        result = dict(row)

        result["existing_duplicate_type"] = None
        result["existing_duplicate_matches"] = []

        if not result.get("valid"):
            annotated_rows.append(
                result
            )
            continue

        data = result.get(
            "data",
            {},
        )

        duplicates = find_possible_client_duplicates(
            cur=cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
            first_name=data.get(
                "first_name",
                "",
            ),
            last_name=data.get(
                "last_name",
                "",
            ),
            phone=data.get(
                "phone",
                "",
            ),
            email=data.get(
                "email",
                "",
            ),
        )

        if duplicates:
            has_strong_match = any(
                duplicate.get(
                    "strong_match"
                )
                for duplicate in duplicates
            )

            result["existing_duplicate_type"] = (
                "strong"
                if has_strong_match
                else "possible"
            )

            result["existing_duplicate_matches"] = (
                duplicates
            )

        annotated_rows.append(
            result
        )

    return annotated_rows



def annotate_peachpos_import_file_duplicates(
    prepared_rows,
):
    """
    Detect repeated processor transaction IDs inside one
    PeachPOS import file.

    Every Import Run belongs to one selected processor, so
    an exact repeated Processor Transaction ID is a strong
    duplicate within that file.
    """
    seen_payment_ids = {}
    annotated_rows = []

    for row in prepared_rows:
        result = dict(row)
        result["duplicate_type"] = None
        result["duplicate_reasons"] = []
        result["duplicate_row_numbers"] = []

        if not result.get("valid"):
            annotated_rows.append(result)
            continue

        data = result.get("data") or {}
        processor_payment_id = str(
            data.get("processor_payment_id") or ""
        ).strip()
        row_number = result.get("row_number")

        if (
            processor_payment_id
            and processor_payment_id in seen_payment_ids
        ):
            result["duplicate_type"] = "strong"
            result["duplicate_reasons"] = [
                "Processor Transaction ID"
            ]
            result["duplicate_row_numbers"] = [
                seen_payment_ids[processor_payment_id]
            ]

        annotated_rows.append(result)

        if processor_payment_id:
            seen_payment_ids.setdefault(
                processor_payment_id,
                row_number,
            )

    return annotated_rows


def annotate_peachpos_import_existing_duplicates(
    cur,
    *,
    spa_id,
    business_unit_id,
    credit_processor_id,
    prepared_rows,
):
    """
    Compare PeachPOS rows against already-posted Income.

    Strong duplicate identity:
      workspace + processor + Processor Transaction ID

    Transaction IDs belonging to a different processor do
    not collide.
    """
    try:
        credit_processor_id = int(credit_processor_id)
    except (TypeError, ValueError):
        raise ImportServiceError(
            "A valid credit processor is required for "
            "PeachPOS duplicate detection."
        )

    payment_ids = sorted({
        str(
            (row.get("data") or {}).get(
                "processor_payment_id"
            )
            or ""
        ).strip()
        for row in prepared_rows
        if row.get("valid")
        and str(
            (row.get("data") or {}).get(
                "processor_payment_id"
            )
            or ""
        ).strip()
    })

    existing_by_payment_id = {}

    if payment_ids:
        cur.execute(
            """
            SELECT
                income_id,
                processor_payment_id
            FROM income
            WHERE spa_id = %s
              AND business_unit_id = %s
              AND income_type = 'PeachPOS'
              AND credit_processor_id = %s
              AND processor_payment_id = ANY(%s)
            """,
            (
                spa_id,
                business_unit_id,
                credit_processor_id,
                payment_ids,
            ),
        )

        for income_id, processor_payment_id in cur.fetchall():
            payment_id = str(
                processor_payment_id or ""
            ).strip()
            if not payment_id:
                continue
            existing_by_payment_id.setdefault(
                payment_id,
                [],
            ).append({
                "income_id": income_id,
                "processor_payment_id": payment_id,
            })

    annotated_rows = []

    for row in prepared_rows:
        result = dict(row)
        result["existing_duplicate_type"] = None
        result["existing_duplicate_matches"] = []

        if not result.get("valid"):
            annotated_rows.append(result)
            continue

        data = result.get("data") or {}
        processor_payment_id = str(
            data.get("processor_payment_id") or ""
        ).strip()

        matches = existing_by_payment_id.get(
            processor_payment_id,
            [],
        )

        if matches:
            result["existing_duplicate_type"] = "strong"
            result["existing_duplicate_matches"] = matches

        annotated_rows.append(result)

    return annotated_rows


# =========================================================
# DURABLE IMPORT RUN STAGING
# =========================================================

def create_import_run(
    uploaded_file,
    *,
    spa_id,
    business_unit_id,
    entity_type,
    requested_by=None,
    allow_repeat=False,
    options=None,
):
    """
    Create one durable Import Engine run and stage its source rows.

    This persists the uploaded data for mapping/review, but does NOT
    create any entity records such as Clients, Income, Expenses,
    Loans, or Appointments.

    Exact-file retry protection prevents an accidental second upload
    of the same file while an earlier run is still actionable or has
    already completed.
    """

    profile = get_import_profile(
        entity_type
    )

    options_payload = dict(
        options or {}
    )

    parsed = parse_import_upload(
        uploaded_file
    )

    uploaded_file.stream.seek(0)
    payload = uploaded_file.stream.read()
    uploaded_file.stream.seek(0)

    source_size_bytes = len(payload)

    if source_size_bytes <= 0:
        raise ImportServiceError(
            "The selected import file is empty."
        )

    source_sha256 = hashlib.sha256(
        payload
    ).hexdigest()

    conn = get_db_connection()
    cur = conn.cursor()

    try:
        workspace_profile = get_workspace_import_profile(
            cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
            entity_type=entity_type,
            options=options_payload,
        )

        mapping = suggest_import_mapping(
            parsed["headers"],
            entity_type=entity_type,
            profile=workspace_profile,
        )

        mapping_payload = {
            "headers": parsed["headers"],
            "suggestions": mapping,
        }

        if not allow_repeat:
            cur.execute(
                """
                SELECT
                    import_run_id,
                    run_status
                FROM import_runs
                WHERE spa_id = %s
                  AND business_unit_id = %s
                  AND entity_type = %s
                  AND source_sha256 = %s
                  AND run_status IN (
                      'mapping',
                      'review',
                      'ready',
                      'importing',
                      'completed'
                  )
                ORDER BY import_run_id DESC
                LIMIT 1
                """,
                (
                    spa_id,
                    business_unit_id,
                    profile["entity_type"],
                    source_sha256,
                ),
            )

            existing = cur.fetchone()

            if existing:
                raise ImportServiceError(
                    "This exact file has already been uploaded "
                    f"as Import Run #{existing[0]} "
                    f"({existing[1]})."
                )

        cur.execute(
            """
            INSERT INTO import_runs (
                spa_id,
                business_unit_id,
                entity_type,
                run_status,
                source_filename,
                source_extension,
                source_size_bytes,
                source_sha256,
                mapping_json,
                options_json,
                total_rows,
                requested_by
            )
            VALUES (
                %s, %s, %s, 'mapping',
                %s, %s, %s, %s, %s, %s, %s, %s
            )
            RETURNING import_run_id
            """,
            (
                spa_id,
                business_unit_id,
                profile["entity_type"],
                parsed["filename"],
                parsed["extension"],
                source_size_bytes,
                source_sha256,
                Json(mapping_payload),
                Json(options_payload),
                parsed["row_count"],
                requested_by,
            ),
        )

        import_run_id = cur.fetchone()[0]

        if parsed["rows"]:
            cur.executemany(
                """
                INSERT INTO import_run_rows (
                    import_run_id,
                    spa_id,
                    business_unit_id,
                    source_row_number,
                    source_data
                )
                VALUES (
                    %s, %s, %s, %s, %s
                )
                """,
                [
                    (
                        import_run_id,
                        spa_id,
                        business_unit_id,
                        source_row["row_number"],
                        Json({
                            "values": source_row["values"],
                        }),
                    )
                    for source_row in parsed["rows"]
                ],
            )

        conn.commit()

        return {
            "import_run_id": import_run_id,
            "entity_type": profile["entity_type"],
            "display_name": profile["display_name"],
            "filename": parsed["filename"],
            "extension": parsed["extension"],
            "source_size_bytes": source_size_bytes,
            "source_sha256": source_sha256,
            "headers": parsed["headers"],
            "mapping": mapping,
            "options": options_payload,
            "row_count": parsed["row_count"],
            "run_status": "mapping",
        }

    except Exception:
        conn.rollback()
        raise

    finally:
        cur.close()
        conn.close()


# =========================================================
# DURABLE IMPORT RUN ANALYSIS / REVIEW
# =========================================================

def analyze_import_run(
    import_run_id,
    *,
    spa_id,
    business_unit_id,
    mapping,
):
    """
    Apply the reviewed mapping to one durable staged Import Run.

    For Clients this:
      - validates required fields and field formats
      - detects duplicates inside the uploaded file
      - detects duplicates against existing PSP clients
      - persists review results per staged row
      - updates Import Run counters/status

    No Client records are created or changed here.
    """

    conn = get_db_connection()
    cur = conn.cursor()

    try:
        cur.execute(
            """
            SELECT
                entity_type,
                run_status,
                mapping_json,
                options_json
            FROM import_runs
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            LIMIT 1
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        run = cur.fetchone()

        if not run:
            raise ImportServiceError(
                "Import Run was not found in this workspace."
            )

        entity_type = str(
            run[0] or ""
        ).strip().lower()

        run_status = str(
            run[1] or ""
        ).strip().lower()

        existing_mapping_json = (
            run[2]
            if isinstance(run[2], dict)
            else {}
        )

        options_payload = (
            run[3]
            if isinstance(run[3], dict)
            else {}
        )

        if run_status in {
            "importing",
            "completed",
            "cancelled",
        }:
            raise ImportServiceError(
                "This Import Run can no longer be remapped."
            )

        profile = get_import_profile(
            entity_type
        )

        valid_target_keys = {
            field["key"]
            for field in profile["fields"]
        }

        required_target_keys = {
            field["key"]
            for field in profile["fields"]
            if field.get("required")
        }

        mapped_target_keys = []

        for item in mapping:
            target_key = str(
                item.get("target_key") or ""
            ).strip()

            if not target_key:
                continue

            if target_key not in valid_target_keys:
                raise ImportServiceError(
                    f"Unknown {profile['display_name']} import "
                    f"field: {target_key}."
                )

            mapped_target_keys.append(
                target_key
            )

        if len(mapped_target_keys) != len(
            set(mapped_target_keys)
        ):
            raise ImportServiceError(
                "More than one source column is mapped to the "
                "same PSP field."
            )

        missing_required = sorted(
            required_target_keys
            - set(mapped_target_keys)
        )

        if missing_required:
            labels_by_key = {
                field["key"]: field["label"]
                for field in profile["fields"]
            }

            missing_labels = [
                labels_by_key.get(
                    key,
                    key,
                )
                for key in missing_required
            ]

            raise ImportServiceError(
                "Required import fields are not mapped: "
                + ", ".join(missing_labels)
                + "."
            )

        cur.execute(
            """
            SELECT
                source_row_number,
                source_data
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            ORDER BY source_row_number
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        staged_rows = cur.fetchall()

        prepared_rows = []

        for (
            source_row_number,
            source_data,
        ) in staged_rows:
            source_data = (
                source_data
                if isinstance(
                    source_data,
                    dict,
                )
                else {}
            )

            values = source_data.get(
                "values",
                []
            )

            mapped = apply_import_mapping(
                values,
                mapping,
                entity_type=entity_type,
            )

            if entity_type == "clients":
                validation = (
                    validate_client_import_row(
                        mapped
                    )
                )
            elif entity_type == "peachpos_income":
                validation = (
                    validate_peachpos_income_import_row(
                        mapped
                    )
                )
            else:
                raise ImportServiceError(
                    f"Unsupported import type: {entity_type}."
                )

            prepared_rows.append({
                "row_number": source_row_number,
                "valid": validation["valid"],
                "data": validation["data"],
                "errors": validation["errors"],
            })

        if entity_type == "clients":
            prepared_rows = (
                validate_client_import_workspace_dropdowns(
                    cur,
                    spa_id=spa_id,
                    business_unit_id=business_unit_id,
                    prepared_rows=prepared_rows,
                )
            )

            prepared_rows = (
                annotate_client_import_file_duplicates(
                    prepared_rows
                )
            )

            prepared_rows = (
                annotate_client_import_existing_duplicates(
                    cur,
                    spa_id=spa_id,
                    business_unit_id=business_unit_id,
                    prepared_rows=prepared_rows,
                )
            )

        elif entity_type == "peachpos_income":
            merchant_verification = (
                verify_peachpos_import_merchant_identity(
                    cur,
                    spa_id=spa_id,
                    business_unit_id=business_unit_id,
                    credit_processor_id=(
                        options_payload.get(
                            "credit_processor_id"
                        )
                    ),
                    prepared_rows=prepared_rows,
                )
            )
            prepared_rows = (
                annotate_peachpos_import_file_duplicates(
                    prepared_rows
                )
            )
            prepared_rows = (
                annotate_peachpos_import_existing_duplicates(
                    cur,
                    spa_id=spa_id,
                    business_unit_id=business_unit_id,
                    credit_processor_id=(
                        options_payload.get(
                            "credit_processor_id"
                        )
                    ),
                    prepared_rows=prepared_rows,
                )
            )
        valid_rows = 0
        invalid_rows = 0
        strong_duplicate_rows = 0
        possible_duplicate_rows = 0

        for row in prepared_rows:
            validation_status = (
                "valid"
                if row.get("valid")
                else "invalid"
            )

            if validation_status == "valid":
                valid_rows += 1
            else:
                invalid_rows += 1

            file_duplicate_type = row.get(
                "duplicate_type"
            )

            existing_duplicate_type = row.get(
                "existing_duplicate_type"
            )

            if (
                file_duplicate_type == "strong"
                or existing_duplicate_type == "strong"
            ):
                duplicate_status = "strong"
                strong_duplicate_rows += 1

            elif (
                file_duplicate_type == "possible"
                or existing_duplicate_type == "possible"
            ):
                duplicate_status = "possible"
                possible_duplicate_rows += 1

            else:
                duplicate_status = "none"

            duplicate_details = {
                "file": {
                    "type": file_duplicate_type,
                    "reasons": row.get(
                        "duplicate_reasons",
                        [],
                    ),
                    "row_numbers": row.get(
                        "duplicate_row_numbers",
                        [],
                    ),
                },
                "existing_psp": {
                    "type": existing_duplicate_type,
                    "matches": row.get(
                        "existing_duplicate_matches",
                        [],
                    ),
                },
            }

            review_decision = None
            if entity_type == "peachpos_income":
                review_decision = (
                    "approved"
                    if validation_status == "valid"
                    and duplicate_status == "none"
                    else "needs_review"
                )

            cur.execute(
                """
                UPDATE import_run_rows
                SET mapped_data = %s,
                    validation_status = %s,
                    validation_errors = %s,
                    duplicate_status = %s,
                    duplicate_details = %s,
                    review_decision = %s,
                    updated_at = CURRENT_TIMESTAMP
                WHERE import_run_id = %s
                  AND spa_id = %s
                  AND business_unit_id = %s
                  AND source_row_number = %s
                """,
                (
                    Json(
                        row.get(
                            "data",
                            {},
                        )
                    ),
                    validation_status,
                    Json(
                        row.get(
                            "errors",
                            [],
                        )
                    ),
                    duplicate_status,
                    Json(
                        duplicate_details
                    ),
                    review_decision,
                    import_run_id,
                    spa_id,
                    business_unit_id,
                    row["row_number"],
                ),
            )

            if cur.rowcount != 1:
                raise ImportServiceError(
                    "A staged import row could not be updated safely."
                )

        needs_review = bool(
            invalid_rows
            or strong_duplicate_rows
            or possible_duplicate_rows
        )

        next_status = (
            "review"
            if needs_review
            else "ready"
        )

        updated_mapping_json = dict(
            existing_mapping_json
        )

        updated_mapping_json[
            "selected"
        ] = mapping

        updated_options_json = dict(
            options_payload
        )
        if entity_type == "peachpos_income":
            updated_options_json[
                "merchant_verification"
            ] = merchant_verification

        cur.execute(
            """
            UPDATE import_runs
            SET mapping_json = %s,
                options_json = %s,
                run_status = %s,
                valid_rows = %s,
                invalid_rows = %s,
                strong_duplicate_rows = %s,
                possible_duplicate_rows = %s,
                last_activity_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                Json(
                    updated_mapping_json
                ),
                Json(
                    updated_options_json
                ),
                next_status,
                valid_rows,
                invalid_rows,
                strong_duplicate_rows,
                possible_duplicate_rows,
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        if cur.rowcount != 1:
            raise ImportServiceError(
                "Import Run could not be updated safely."
            )

        conn.commit()

        return {
            "import_run_id": import_run_id,
            "entity_type": entity_type,
            "run_status": next_status,
            "total_rows": len(
                prepared_rows
            ),
            "valid_rows": valid_rows,
            "invalid_rows": invalid_rows,
            "strong_duplicate_rows": (
                strong_duplicate_rows
            ),
            "possible_duplicate_rows": (
                possible_duplicate_rows
            ),
        }

    except Exception:
        conn.rollback()
        raise

    finally:
        cur.close()
        conn.close()


# =========================================================
# IMPORT RUN RECORDS
# =========================================================

def get_import_run_records(
    import_run_id,
    *,
    spa_id,
    business_unit_id,
):
    """
    Load the complete durable transaction/review view for one
    Import Run.

    Returns the batch header and every staged source row with:
      - original source data
      - mapped data
      - validation state/errors
      - duplicate state/details
      - review decision
      - import/posting state
      - resulting Income record ID when posted

    This is read-only and strictly workspace scoped.
    """
    conn = get_db_connection()
    cur = conn.cursor()

    try:
        cur.execute(
            """
            SELECT
                entity_type,
                run_status,
                source_filename,
                source_extension,
                source_size_bytes,
                options_json,
                total_rows,
                valid_rows,
                invalid_rows,
                strong_duplicate_rows,
                possible_duplicate_rows,
                imported_rows,
                skipped_rows,
                error_rows,
                requested_at,
                started_at,
                completed_at,
                failure_message
            FROM import_runs
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            LIMIT 1
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        run = cur.fetchone()

        if not run:
            raise ImportServiceError(
                "Import Run was not found in this workspace."
            )

        entity_type = str(
            run[0] or ""
        ).strip().lower()

        options_payload = (
            run[5]
            if isinstance(run[5], dict)
            else {}
        )

        profile = get_workspace_import_profile(
            cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
            entity_type=entity_type,
            options=options_payload,
        )

        credit_processor_name = None

        if entity_type == "peachpos_income":
            credit_processor_id = options_payload.get(
                "credit_processor_id"
            )

            try:
                credit_processor_id = int(
                    credit_processor_id
                )
            except (TypeError, ValueError):
                credit_processor_id = None

            if credit_processor_id:
                cur.execute(
                    """
                    SELECT credit_processor_name
                    FROM credit_processors
                    WHERE credit_processor_id = %s
                      AND spa_id = %s
                      AND business_unit_id = %s
                    LIMIT 1
                    """,
                    (
                        credit_processor_id,
                        spa_id,
                        business_unit_id,
                    ),
                )
                processor_row = cur.fetchone()

                if processor_row:
                    credit_processor_name = (
                        str(
                            processor_row[0] or ""
                        ).strip()
                        or None
                    )

        cur.execute(
            """
            SELECT
                import_run_row_id,
                source_row_number,
                source_data,
                mapped_data,
                validation_status,
                validation_errors,
                duplicate_status,
                duplicate_details,
                review_decision,
                import_status,
                imported_record_id,
                error_message
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            ORDER BY source_row_number
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        records = []

        for (
            import_run_row_id,
            source_row_number,
            source_data,
            mapped_data,
            validation_status,
            validation_errors,
            duplicate_status,
            duplicate_details,
            review_decision,
            import_status,
            imported_record_id,
            error_message,
        ) in cur.fetchall():

            source_data = (
                source_data
                if isinstance(source_data, dict)
                else {}
            )

            mapped_data = (
                mapped_data
                if isinstance(mapped_data, dict)
                else {}
            )

            validation_errors = (
                validation_errors
                if isinstance(validation_errors, list)
                else []
            )

            duplicate_details = (
                duplicate_details
                if isinstance(duplicate_details, dict)
                else {}
            )

            records.append({
                "import_run_row_id": import_run_row_id,
                "source_row_number": source_row_number,
                "source_values": source_data.get(
                    "values",
                    [],
                ),
                "mapped_data": mapped_data,
                "validation_status": validation_status,
                "validation_errors": validation_errors,
                "duplicate_status": duplicate_status,
                "duplicate_details": duplicate_details,
                "review_decision": review_decision,
                "import_status": import_status,
                "imported_record_id": imported_record_id,
                "error_message": error_message,
            })

        approved_rows = sum(
            1
            for record in records
            if record.get("review_decision") == "approved"
        )

        needs_review_rows = sum(
            1
            for record in records
            if record.get("review_decision") == "needs_review"
        )

        hold_rows = sum(
            1
            for record in records
            if record.get("review_decision") == "hold"
        )

        discard_rows = sum(
            1
            for record in records
            if record.get("review_decision") == "discard"
        )

        financial_summary = {
            "transactions": 0,
            "sales": Decimal("0.00"),
            "tax": Decimal("0.00"),
            "tips": Decimal("0.00"),
            "gross": Decimal("0.00"),
            "processor_fees": Decimal("0.00"),
            "net_received": Decimal("0.00"),
        }

        financial_field_map = (
            ("pos_amount", "sales"),
            ("tax_amount", "tax"),
            ("tip_amount", "tips"),
            ("total_amount", "gross"),
            (
                "processing_fee_amount",
                "processor_fees",
            ),
            ("net_received", "net_received"),
        )

        for record in records:
            if (
                record.get("review_decision")
                != "approved"
            ):
                continue

            financial_summary[
                "transactions"
            ] += 1

            data = record.get("mapped_data") or {}

            for field_key, summary_key in (
                financial_field_map
            ):
                raw_value = str(
                    data.get(field_key) or ""
                ).strip()

                if not raw_value:
                    continue

                try:
                    amount = Decimal(
                        raw_value
                    ).quantize(
                        Decimal("0.01")
                    )
                except (
                    InvalidOperation,
                    ValueError,
                ):
                    continue

                financial_summary[
                    summary_key
                ] += amount

        total_gross = financial_summary["gross"]
        total_net = financial_summary[
            "net_received"
        ]

        return {
            "import_run_id": import_run_id,
            "entity_type": entity_type,
            "display_name": profile["display_name"],
            "run_status": run[1],
            "filename": run[2],
            "extension": run[3],
            "source_size_bytes": run[4],
            "options": options_payload,
            "credit_processor_name": credit_processor_name,
            "fields": [
                {
                    "key": field["key"],
                    "label": field["label"],
                    "required": bool(
                        field.get("required")
                    ),
                }
                for field in profile["fields"]
            ],
            "total_rows": run[6],
            "valid_rows": run[7],
            "approved_rows": approved_rows,
            "needs_review_rows": needs_review_rows,
            "hold_rows": hold_rows,
            "discard_rows": discard_rows,
            "financial_summary": financial_summary,
            "total_gross": total_gross,
            "total_net": total_net,
            "invalid_rows": run[8],
            "strong_duplicate_rows": run[9],
            "possible_duplicate_rows": run[10],
            "imported_rows": run[11],
            "skipped_rows": run[12],
            "error_rows": run[13],
            "requested_at": run[14],
            "started_at": run[15],
            "completed_at": run[16],
            "failure_message": run[17],
            "records": records,
        }

    finally:
        cur.close()
        conn.close()


# =========================================================
# IMPORT RUN RECORD DETAILS
# =========================================================

def get_import_run_record_details(
    import_run_id,
    import_run_row_id,
    *,
    spa_id,
    business_unit_id,
):
    """
    Load one durable PeachPOS Import Run transaction for a
    read-only transaction-details view.

    Returns standard imported transaction data plus all 12
    processor-specific fields using the processor's current
    workspace-configured display labels.

    No writes are performed.
    """
    conn = get_db_connection()
    cur = conn.cursor()

    try:
        cur.execute(
            """
            SELECT
                r.entity_type,
                r.options_json,
                rr.import_run_row_id,
                rr.source_row_number,
                rr.source_data,
                rr.mapped_data,
                rr.validation_status,
                rr.validation_errors,
                rr.duplicate_status,
                rr.duplicate_details,
                rr.review_decision,
                rr.import_status,
                rr.imported_record_id,
                rr.error_message
            FROM import_runs r
            JOIN import_run_rows rr
              ON rr.import_run_id = r.import_run_id
             AND rr.spa_id = r.spa_id
             AND rr.business_unit_id = r.business_unit_id
            WHERE r.import_run_id = %s
              AND r.spa_id = %s
              AND r.business_unit_id = %s
              AND rr.import_run_row_id = %s
            LIMIT 1
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
                import_run_row_id,
            ),
        )

        row = cur.fetchone()

        if not row:
            raise ImportServiceError(
                "Import Record was not found in this workspace."
            )

        entity_type = str(
            row[0] or ""
        ).strip().lower()

        if entity_type != "peachpos_income":
            raise ImportServiceError(
                "This transaction details view is available "
                "for PeachPOS Income imports only."
            )

        options_payload = (
            row[1]
            if isinstance(row[1], dict)
            else {}
        )

        credit_processor_id = options_payload.get(
            "credit_processor_id"
        )

        try:
            credit_processor_id = int(
                credit_processor_id
            )
        except (TypeError, ValueError):
            raise ImportServiceError(
                "This PeachPOS Import Run does not have a valid "
                "credit processor."
            )

        processor_labels = (
            get_peachpos_processor_field_labels(
                cur,
                spa_id=spa_id,
                business_unit_id=business_unit_id,
                credit_processor_id=credit_processor_id,
            )
        )

        processor_definitions = (
            get_peachpos_processor_field_definitions()
        )

        source_data = (
            row[4]
            if isinstance(row[4], dict)
            else {}
        )

        mapped_data = (
            row[5]
            if isinstance(row[5], dict)
            else {}
        )

        validation_errors = (
            row[7]
            if isinstance(row[7], list)
            else []
        )

        duplicate_details = (
            row[9]
            if isinstance(row[9], dict)
            else {}
        )

        processor_fields = []

        for definition in processor_definitions:
            field_key = definition["field_key"]

            processor_fields.append({
                "field_key": field_key,
                "label": (
                    processor_labels.get(field_key)
                    or definition["default_label"]
                ),
                "default_label": definition[
                    "default_label"
                ],
                "value": mapped_data.get(
                    field_key
                ),
            })

        return {
            "import_run_id": import_run_id,
            "import_run_row_id": row[2],
            "source_row_number": row[3],
            "source_values": source_data.get(
                "values",
                [],
            ),
            "mapped_data": mapped_data,
            "validation_status": row[6],
            "validation_errors": validation_errors,
            "duplicate_status": row[8],
            "duplicate_details": duplicate_details,
            "review_decision": row[10],
            "import_status": row[11],
            "imported_record_id": row[12],
            "error_message": row[13],
            "options": options_payload,
            "credit_processor_id": credit_processor_id,
            "processor_fields": processor_fields,
        }

    finally:
        cur.close()
        conn.close()


# =========================================================
# CLIENT IMPORT EXECUTION
# =========================================================

PEACHPOS_INCOME_IMPORT_PACKET_SIZE = 50
PEACHPOS_INCOME_IMPORT_PACKET_MAX_ROWS = 50

CLIENT_IMPORT_PACKET_SIZE = 50
CLIENT_IMPORT_PACKET_MAX_ROWS = 50


def process_client_import_packet(
    import_run_id,
    *,
    spa_id,
    business_unit_id,
    packet_size=CLIENT_IMPORT_PACKET_SIZE,
):
    """
    Import one bounded packet of validated Client rows.

    Safety rules:
      - only READY / already-IMPORTING Client runs may execute
      - only valid, non-duplicate, still-pending rows are imported
      - each Client insert and row-state update are atomic
      - communication permissions are explicitly conservative
      - no Square synchronization occurs here
      - repeated calls are safe; imported rows are never reinserted
    """

    try:
        packet_size = int(packet_size)
    except (TypeError, ValueError):
        raise ImportServiceError(
            "Import packet size must be a whole number."
        )

    if (
        packet_size < 1
        or packet_size > CLIENT_IMPORT_PACKET_MAX_ROWS
    ):
        raise ImportServiceError(
            "Import packet size must be between 1 and "
            f"{CLIENT_IMPORT_PACKET_MAX_ROWS}."
        )

    conn = get_db_connection()
    cur = conn.cursor()

    try:
        # Lock the run so two requests cannot process the same
        # packet at the same time.
        cur.execute(
            """
            SELECT
                entity_type,
                run_status
            FROM import_runs
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            FOR UPDATE
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        run = cur.fetchone()

        if not run:
            raise ImportServiceError(
                "Import Run was not found in this workspace."
            )

        entity_type = str(
            run[0] or ""
        ).strip().lower()

        run_status = str(
            run[1] or ""
        ).strip().lower()

        if entity_type != "clients":
            raise ImportServiceError(
                "This importer currently executes Client runs only."
            )

        # A completed run is safely idempotent.
        if run_status == "completed":
            cur.execute(
                """
                SELECT
                    total_rows,
                    imported_rows,
                    skipped_rows,
                    error_rows
                FROM import_runs
                WHERE import_run_id = %s
                """,
                (import_run_id,),
            )

            counts = cur.fetchone()

            conn.rollback()

            return {
                "import_run_id": import_run_id,
                "run_status": "completed",
                "total_rows": counts[0],
                "imported_rows": counts[1],
                "skipped_rows": counts[2],
                "error_rows": counts[3],
                "processed_this_packet": 0,
                "remaining_rows": 0,
            }

        if run_status not in {
            "ready",
            "importing",
        }:
            raise ImportServiceError(
                "This Import Run is not ready to import. "
                "Complete its review first."
            )

        # READY must mean every still-pending row is fully valid
        # and has no unresolved duplicate warning.
        cur.execute(
            """
            SELECT COUNT(*)
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND import_status = 'pending'
              AND (
                  validation_status <> 'valid'
                  OR duplicate_status <> 'none'
              )
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        unresolved_rows = cur.fetchone()[0]

        if unresolved_rows:
            raise ImportServiceError(
                "This Import Run still has rows requiring review."
            )

        cur.execute(
            """
            UPDATE import_runs
            SET run_status = 'importing',
                started_at = COALESCE(
                    started_at,
                    CURRENT_TIMESTAMP
                ),
                last_activity_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        cur.execute(
            """
            SELECT
                import_run_row_id,
                source_row_number,
                mapped_data
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND validation_status = 'valid'
              AND duplicate_status = 'none'
              AND import_status = 'pending'
            ORDER BY source_row_number
            LIMIT %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
                packet_size,
            ),
        )

        packet_rows = cur.fetchall()
        processed_this_packet = 0

        for (
            import_run_row_id,
            source_row_number,
            mapped_data,
        ) in packet_rows:

            data = (
                mapped_data
                if isinstance(mapped_data, dict)
                else {}
            )

            cur.execute(
                "SAVEPOINT client_import_row"
            )

            try:
                cur.execute(
                    """
                    INSERT INTO clients (
                        spa_id,
                        business_unit_id,
                        first_name,
                        last_name,
                        phone,
                        email,
                        birth_date,
                        address,
                        city,
                        state,
                        zip,
                        client_status,
                        preferred_language,
                        ok_to_call,
                        ok_to_text,
                        ok_to_email,
                        preferred_contact_method,
                        emergency_contact_name,
                        emergency_contact_phone,
                        referred_by,
                        notes_one,
                        notes_two,
                        notes_three,
                        active_client,
                        sms_opt_in,
                        sms_opt_out,
                        email_opt_in,
                        email_opt_out,
                        sms_marketing_status,
                        email_marketing_status,
                        sms_consent_timestamp,
                        email_consent_timestamp
                    )
                    VALUES (
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s
                    )
                    RETURNING client_id
                    """,
                    (
                        spa_id,
                        business_unit_id,
                        data.get("first_name", ""),
                        data.get("last_name", ""),
                        data.get("phone", ""),
                        data.get("email", ""),
                        data.get("birth_date") or None,
                        data.get("address", ""),
                        data.get("city", ""),
                        data.get("state", ""),
                        data.get("zip", ""),
                        data.get(
                            "client_status",
                            "Current",
                        ),
                        data.get(
                            "preferred_language"
                        ) or None,
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "ok_to_call"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "ok_to_text"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "ok_to_email"
                        ],
                        data.get(
                            "preferred_contact_method"
                        ) or None,
                        data.get(
                            "emergency_contact_name",
                            "",
                        ),
                        data.get(
                            "emergency_contact_phone",
                            "",
                        ),
                        None,
                        data.get("notes_one", ""),
                        data.get("notes_two", ""),
                        data.get("notes_three", ""),
                        bool(
                            data.get(
                                "active_client",
                                True,
                            )
                        ),
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "sms_opt_in"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "sms_opt_out"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "email_opt_in"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "email_opt_out"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "sms_marketing_status"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "email_marketing_status"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "sms_consent_timestamp"
                        ],
                        CLIENT_IMPORT_COMMUNICATION_DEFAULTS[
                            "email_consent_timestamp"
                        ],
                    ),
                )

                client_id = cur.fetchone()[0]

                client_information_keys = tuple(
                    field["key"]
                    for field in CLIENT_INFORMATION_IMPORT_FIELDS
                )

                has_client_information_data = any(
                    (
                        data.get(field_key) is not None
                        and (
                            not isinstance(
                                data.get(field_key),
                                str,
                            )
                            or bool(
                                data.get(field_key).strip()
                            )
                        )
                    )
                    for field_key in client_information_keys
                    if field_key in data
                )

                if has_client_information_data:
                    cur.execute(
                        """
                        INSERT INTO client_health_profile (
                            spa_id,
                            client_id,
                            skin_type_id,
                            fitzpatrick_id,
                            skin_concerns,
                            skin_conditions,
                            allergies,
                            medications,
                            current_medical_conditions,
                            past_medical_treatments,
                            recent_injections,
                            recent_laser,
                            pregnant,
                            nursing,
                            using_retinol,
                            using_accutane,
                            sun_exposure_level,
                            last_facial_date,
                            notes1,
                            notes2,
                            notes3
                        )
                        VALUES (
                            %s, %s, %s, %s, %s, %s, %s,
                            %s, %s, %s, %s, %s, %s, %s,
                            %s, %s, %s, %s, %s, %s, %s
                        )
                        """,
                        (
                            spa_id,
                            client_id,
                            data.get("skin_type_id") or None,
                            data.get("fitzpatrick_id") or None,
                            data.get("skin_concerns") or None,
                            data.get("skin_conditions") or None,
                            data.get("allergies") or None,
                            data.get("medications") or None,
                            data.get(
                                "current_medical_conditions"
                            ) or None,
                            data.get(
                                "past_medical_treatments"
                            ) or None,
                            data.get("recent_injections"),
                            data.get("recent_laser"),
                            data.get("pregnant"),
                            data.get("nursing"),
                            data.get("using_retinol"),
                            data.get("using_accutane"),
                            data.get("sun_exposure_level") or None,
                            data.get("last_facial_date") or None,
                            data.get("notes1") or None,
                            data.get("notes2") or None,
                            data.get("notes3") or None,
                        ),
                    )

                cur.execute(
                    """
                    UPDATE import_run_rows
                    SET import_status = 'imported',
                        imported_record_id = %s,
                        error_message = NULL,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE import_run_row_id = %s
                      AND import_run_id = %s
                      AND spa_id = %s
                      AND business_unit_id = %s
                      AND import_status = 'pending'
                    """,
                    (
                        client_id,
                        import_run_row_id,
                        import_run_id,
                        spa_id,
                        business_unit_id,
                    ),
                )

                if cur.rowcount != 1:
                    raise ImportServiceError(
                        "The staged Client row changed while "
                        "it was being imported."
                    )

                cur.execute(
                    "RELEASE SAVEPOINT client_import_row"
                )

                processed_this_packet += 1

            except Exception as exc:
                cur.execute(
                    "ROLLBACK TO SAVEPOINT client_import_row"
                )

                cur.execute(
                    "RELEASE SAVEPOINT client_import_row"
                )

                cur.execute(
                    """
                    UPDATE import_run_rows
                    SET import_status = 'error',
                        imported_record_id = NULL,
                        error_message = %s,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE import_run_row_id = %s
                      AND import_run_id = %s
                      AND spa_id = %s
                      AND business_unit_id = %s
                    """,
                    (
                        str(exc)[:1000],
                        import_run_row_id,
                        import_run_id,
                        spa_id,
                        business_unit_id,
                    ),
                )

        # Recalculate from durable row state rather than relying
        # on in-memory increments. This keeps retries deterministic.
        cur.execute(
            """
            SELECT
                COUNT(*) FILTER (
                    WHERE import_status = 'imported'
                ),
                COUNT(*) FILTER (
                    WHERE import_status = 'skipped'
                ),
                COUNT(*) FILTER (
                    WHERE import_status = 'error'
                ),
                COUNT(*) FILTER (
                    WHERE import_status = 'pending'
                )
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        (
            imported_rows,
            skipped_rows,
            error_rows,
            remaining_rows,
        ) = cur.fetchone()

        if remaining_rows > 0:
            next_status = "importing"
            failure_message = None
            completed_at_sql = "completed_at"

        elif error_rows > 0:
            next_status = "failed"
            failure_message = (
                "One or more Client rows could not be imported."
            )
            completed_at_sql = "CURRENT_TIMESTAMP"

        else:
            next_status = "completed"
            failure_message = None
            completed_at_sql = "CURRENT_TIMESTAMP"

        cur.execute(
            f"""
            UPDATE import_runs
            SET run_status = %s,
                imported_rows = %s,
                skipped_rows = %s,
                error_rows = %s,
                failure_message = %s,
                completed_at = {completed_at_sql},
                last_activity_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                next_status,
                imported_rows,
                skipped_rows,
                error_rows,
                failure_message,
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        conn.commit()

        return {
            "import_run_id": import_run_id,
            "run_status": next_status,
            "imported_rows": imported_rows,
            "skipped_rows": skipped_rows,
            "error_rows": error_rows,
            "processed_this_packet": (
                processed_this_packet
            ),
            "remaining_rows": remaining_rows,
        }

    except Exception:
        conn.rollback()
        raise

    finally:
        cur.close()
        conn.close()



# =========================================================
# PEACHPOS INCOME IMPORT POSTING
# =========================================================

def process_peachpos_income_import_packet(
    import_run_id,
    *,
    spa_id,
    business_unit_id,
    packet_size=PEACHPOS_INCOME_IMPORT_PACKET_SIZE,
):
    """
    Post one bounded packet of approved PeachPOS Income rows.

    Safety rules:
      - only READY / already-IMPORTING PeachPOS Income runs execute
      - only valid, non-duplicate, approved, pending rows post
      - held / needs-review rows remain staged and unposted
      - discarded rows are marked skipped and never post
      - each Income insert and staged-row update are atomic
      - Processor Fields 1-8 are the active imported custom fields
      - repeated calls are safe; imported rows are never reinserted
      - database uniqueness protects processor transaction identity
    """
    try:
        packet_size = int(packet_size)
    except (TypeError, ValueError):
        raise ImportServiceError(
            "Import packet size must be a whole number."
        )

    if (
        packet_size < 1
        or packet_size > PEACHPOS_INCOME_IMPORT_PACKET_MAX_ROWS
    ):
        raise ImportServiceError(
            "Import packet size must be between 1 and "
            f"{PEACHPOS_INCOME_IMPORT_PACKET_MAX_ROWS}."
        )

    conn = get_db_connection()
    cur = conn.cursor()

    try:
        # Lock the run so two requests cannot post the same packet
        # at the same time.
        cur.execute(
            """
            SELECT
                entity_type,
                run_status,
                options_json
            FROM import_runs
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            FOR UPDATE
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )
        run = cur.fetchone()

        if not run:
            raise ImportServiceError(
                "Import Run was not found in this workspace."
            )

        entity_type = str(run[0] or "").strip().lower()
        run_status = str(run[1] or "").strip().lower()
        options = (
            run[2]
            if isinstance(run[2], dict)
            else {}
        )

        if entity_type != "peachpos_income":
            raise ImportServiceError(
                "This importer executes PeachPOS Income runs only."
            )

        if run_status == "completed":
            cur.execute(
                """
                SELECT
                    total_rows,
                    imported_rows,
                    skipped_rows,
                    error_rows
                FROM import_runs
                WHERE import_run_id = %s
                """,
                (import_run_id,),
            )
            counts = cur.fetchone()
            conn.rollback()

            return {
                "import_run_id": import_run_id,
                "run_status": "completed",
                "total_rows": counts[0],
                "imported_rows": counts[1],
                "skipped_rows": counts[2],
                "error_rows": counts[3],
                "processed_this_packet": 0,
                "remaining_rows": 0,
            }

        if run_status not in {
            "ready",
            "importing",
        }:
            raise ImportServiceError(
                "This PeachPOS Import Run is not ready to post."
            )

        credit_processor_id = options.get(
            "credit_processor_id"
        )
        try:
            credit_processor_id = int(
                credit_processor_id
            )
        except (TypeError, ValueError):
            raise ImportServiceError(
                "This PeachPOS Import Run does not have a valid "
                "credit processor."
            )

        event_name = str(
            options.get("event_name") or ""
        ).strip() or None

        # Verify that the selected processor belongs to this exact
        # Provider Workspace. Historical/inactive processors remain
        # valid for durable import history.
        cur.execute(
            """
            SELECT
                credit_processor_name,
                percentage_fee,
                flat_fee,
                additional_fee,
                merchant_account_identifier
            FROM credit_processors
            WHERE credit_processor_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                credit_processor_id,
                spa_id,
                business_unit_id,
            ),
        )
        processor = cur.fetchone()

        if not processor:
            raise ImportServiceError(
                "The selected credit processor was not found in "
                "this workspace."
            )

        processor_name = str(
            processor[0] or ""
        ).strip()

        if not processor_name:
            raise ImportServiceError(
                "The selected credit processor does not have a "
                "valid name."
            )

        # -------------------------------------------------
        # Defense-in-depth Merchant ID posting gate.
        #
        # Analysis must have already verified one Merchant ID
        # from the uploaded file against Processor Company Setup.
        # Recheck that durable verification snapshot here before
        # any Income row can be created.
        # -------------------------------------------------

        if processor_name.lower() == "square":
            raise ImportServiceError(
                "Square transactions use PSP's native Square "
                "integration and cannot be posted through the "
                "generic PeachPOS Income Import."
            )

        merchant_verification = options.get(
            "merchant_verification"
        )
        if not isinstance(
            merchant_verification,
            dict,
        ):
            raise ImportServiceError(
                "This PeachPOS Import Run does not have a "
                "Merchant ID / Account ID verification. "
                "Re-analyze the import before posting."
            )

        if merchant_verification.get("verified") is not True:
            raise ImportServiceError(
                "Merchant ID / Account ID verification was not "
                "successful. Import posting is blocked."
            )

        detected_identifier = str(
            merchant_verification.get(
                "detected_identifier"
            )
            or ""
        ).strip()

        verified_configured_identifier = str(
            merchant_verification.get(
                "configured_identifier_at_verification"
            )
            or ""
        ).strip()

        verified_processor_company = str(
            merchant_verification.get(
                "processor_company"
            )
            or ""
        ).strip()

        current_configured_identifier = str(
            processor[4] or ""
        ).strip()

        if (
            not detected_identifier
            or not verified_configured_identifier
            or detected_identifier
                != verified_configured_identifier
        ):
            raise ImportServiceError(
                "This PeachPOS Import Run has an invalid "
                "Merchant ID / Account ID verification record. "
                "Posting is blocked."
            )

        if (
            verified_processor_company
            and verified_processor_company != processor_name
        ):
            raise ImportServiceError(
                "The processor company changed after Merchant "
                "ID verification. Re-analyze the import before "
                "posting."
            )

        if not current_configured_identifier:
            raise ImportServiceError(
                "Merchant ID / Account ID is no longer "
                "configured for this processor. "
                "Posting is blocked."
            )

        if (
            current_configured_identifier
            != verified_configured_identifier
        ):
            raise ImportServiceError(
                "Merchant ID / Account ID in Processor Company "
                "Setup changed after this import was verified. "
                "Re-analyze the import before posting."
            )

        processor_percentage_fee = (
            processor[1]
            if processor[1] is not None
            else Decimal("0")
        )
        processor_flat_fee = (
            processor[2]
            if processor[2] is not None
            else Decimal("0")
        )
        processor_additional_fee = (
            processor[3]
            if processor[3] is not None
            else Decimal("0")
        )

        # A row cannot be posted merely because somebody marked it
        # approved. It must also have passed validation and duplicate
        # analysis.
        cur.execute(
            """
            SELECT COUNT(*)
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND import_status = 'pending'
              AND review_decision = 'approved'
              AND (
                    validation_status <> 'valid'
                    OR duplicate_status <> 'none'
              )
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )
        unsafe_approved_rows = cur.fetchone()[0]

        if unsafe_approved_rows:
            raise ImportServiceError(
                "One or more approved PeachPOS rows still require "
                "validation or duplicate review."
            )

        # Discard is a terminal review decision. Mark those rows
        # skipped so they remain in the Import Record but can never
        # accidentally post later.
        cur.execute(
            """
            UPDATE import_run_rows
            SET import_status = 'skipped',
                imported_record_id = NULL,
                error_message = NULL,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND import_status = 'pending'
              AND review_decision = 'discard'
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        cur.execute(
            """
            UPDATE import_runs
            SET run_status = 'importing',
                started_at = COALESCE(
                    started_at,
                    CURRENT_TIMESTAMP
                ),
                last_activity_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        cur.execute(
            """
            SELECT
                import_run_row_id,
                source_row_number,
                mapped_data
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND validation_status = 'valid'
              AND duplicate_status = 'none'
              AND review_decision = 'approved'
              AND import_status = 'pending'
            ORDER BY source_row_number
            LIMIT %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
                packet_size,
            ),
        )
        packet_rows = cur.fetchall()
        processed_this_packet = 0

        for (
            import_run_row_id,
            source_row_number,
            mapped_data,
        ) in packet_rows:
            data = (
                mapped_data
                if isinstance(mapped_data, dict)
                else {}
            )

            cur.execute(
                "SAVEPOINT peachpos_income_import_row"
            )

            try:
                transaction_at_raw = str(
                    data.get("transaction_at") or ""
                ).strip()

                if not transaction_at_raw:
                    raise ImportServiceError(
                        "Transaction Date/Time is required."
                    )

                try:
                    transaction_at = datetime.fromisoformat(
                        transaction_at_raw
                    )
                except ValueError:
                    raise ImportServiceError(
                        "Transaction Date/Time is invalid."
                    )

                processor_payment_id = str(
                    data.get("processor_payment_id") or ""
                ).strip()

                if not processor_payment_id:
                    raise ImportServiceError(
                        "Processor Transaction ID is required."
                    )

                def money_value(key, *, required=False):
                    raw = data.get(key)

                    if raw is None or str(raw).strip() == "":
                        if required:
                            raise ImportServiceError(
                                f"{key} is required."
                            )
                        return Decimal("0.00")

                    try:
                        return Decimal(str(raw))
                    except (InvalidOperation, ValueError):
                        raise ImportServiceError(
                            f"{key} is not a valid money amount."
                        )

                pos_amount = money_value(
                    "pos_amount",
                    required=True,
                )
                tax_amount = money_value("tax_amount")
                tip_amount = money_value("tip_amount")
                total_amount = money_value("total_amount")
                processing_fee_amount = money_value(
                    "processing_fee_amount"
                )
                net_received = money_value("net_received")

                # Income accounting columns are historically NOT NULL.
                # Preserve whether each optional amount was actually
                # supplied by the processor so a compatibility 0.00 is
                # never mistaken for processor-reported zero.
                optional_money_fields = (
                    "tax_amount",
                    "tip_amount",
                    "total_amount",
                    "processing_fee_amount",
                    "net_received",
                )
                processor_import_data = {
                    "source_presence": {
                        key: (
                            data.get(key) is not None
                            and str(data.get(key)).strip() != ""
                        )
                        for key in optional_money_fields
                    }
                }

                payment_method = str(
                    data.get("payment_method") or ""
                ).strip() or None

                processor_status = str(
                    data.get("processor_status") or ""
                ).strip() or None

                processor_fields = [
                    (
                        str(
                            data.get(
                                f"processor_field_{number}"
                            )
                            or ""
                        ).strip()
                        or None
                    )
                    for number in range(1, 9)
                ]

                description = (
                    f"PeachPOS {processor_name} Sale"
                )

                cur.execute(
                    """
                    INSERT INTO income (
                        income_date,
                        income_type,
                        description,
                        service_amount,
                        retail_amount,
                        pos_amount,
                        tax_amount,
                        tip_amount,
                        total_amount,
                        payment_method,
                        processor_payment_id,
                        spa_id,
                        credit_processor_id,
                        processing_fee_amount,
                        net_received,
                        processor_percentage_fee,
                        processor_flat_fee,
                        processor_additional_fee,
                        business_unit_id,
                        transaction_at,
                        event_name,
                        processor_status,
                        processor_import_data,
                        processor_field_1,
                        processor_field_2,
                        processor_field_3,
                        processor_field_4,
                        processor_field_5,
                        processor_field_6,
                        processor_field_7,
                        processor_field_8
                    )
                    VALUES (
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s, %s, %s, %s, %s,
                        %s
                    )
                    RETURNING income_id
                    """,
                    (
                        transaction_at.date(),
                        "PeachPOS",
                        description,
                        Decimal("0.00"),
                        Decimal("0.00"),
                        pos_amount,
                        tax_amount,
                        tip_amount,
                        total_amount,
                        payment_method,
                        processor_payment_id,
                        spa_id,
                        credit_processor_id,
                        processing_fee_amount,
                        net_received,
                        processor_percentage_fee,
                        processor_flat_fee,
                        processor_additional_fee,
                        business_unit_id,
                        transaction_at,
                        event_name,
                        processor_status,
                        Json(processor_import_data),
                        *processor_fields,
                    ),
                )
                income_id = cur.fetchone()[0]

                cur.execute(
                    """
                    UPDATE import_run_rows
                    SET import_status = 'imported',
                        imported_record_id = %s,
                        error_message = NULL,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE import_run_row_id = %s
                      AND import_run_id = %s
                      AND spa_id = %s
                      AND business_unit_id = %s
                      AND review_decision = 'approved'
                      AND import_status = 'pending'
                    """,
                    (
                        income_id,
                        import_run_row_id,
                        import_run_id,
                        spa_id,
                        business_unit_id,
                    ),
                )

                if cur.rowcount != 1:
                    raise ImportServiceError(
                        "The staged PeachPOS row changed while "
                        "it was being posted."
                    )

                cur.execute(
                    "RELEASE SAVEPOINT "
                    "peachpos_income_import_row"
                )
                processed_this_packet += 1

            except Exception as exc:
                cur.execute(
                    "ROLLBACK TO SAVEPOINT "
                    "peachpos_income_import_row"
                )
                cur.execute(
                    "RELEASE SAVEPOINT "
                    "peachpos_income_import_row"
                )

                cur.execute(
                    """
                    UPDATE import_run_rows
                    SET import_status = 'error',
                        imported_record_id = NULL,
                        error_message = %s,
                        updated_at = CURRENT_TIMESTAMP
                    WHERE import_run_row_id = %s
                      AND import_run_id = %s
                      AND spa_id = %s
                      AND business_unit_id = %s
                    """,
                    (
                        str(exc)[:1000],
                        import_run_row_id,
                        import_run_id,
                        spa_id,
                        business_unit_id,
                    ),
                )

        # Durable counters are always recalculated from staged row
        # state so retries remain deterministic.
        cur.execute(
            """
            SELECT
                COUNT(*) FILTER (
                    WHERE import_status = 'imported'
                ),
                COUNT(*) FILTER (
                    WHERE import_status = 'skipped'
                ),
                COUNT(*) FILTER (
                    WHERE import_status = 'error'
                )
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )
        (
            imported_rows,
            skipped_rows,
            error_rows,
        ) = cur.fetchone()

        # Only approved, still-pending rows are remaining posting
        # work. Hold and needs_review deliberately remain staged.
        cur.execute(
            """
            SELECT COUNT(*)
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND review_decision = 'approved'
              AND import_status = 'pending'
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )
        remaining_rows = cur.fetchone()[0]

        # Hold and needs_review rows are intentionally unresolved.
        # Keep the run open so they can be reviewed, approved, and
        # posted later instead of being stranded by a completed run.
        cur.execute(
            """
            SELECT COUNT(*)
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
              AND review_decision IN ('hold', 'needs_review')
              AND import_status = 'pending'
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )
        unresolved_review_rows = cur.fetchone()[0]

        if remaining_rows > 0:
            next_status = "importing"
            failure_message = None
            completed_at_sql = "completed_at"
        elif error_rows > 0:
            next_status = "failed"
            failure_message = (
                "One or more PeachPOS transactions could not "
                "be posted."
            )
            completed_at_sql = "CURRENT_TIMESTAMP"
        elif unresolved_review_rows > 0:
            next_status = "ready"
            failure_message = None
            completed_at_sql = "NULL"
        else:
            next_status = "completed"
            failure_message = None
            completed_at_sql = "CURRENT_TIMESTAMP"

        cur.execute(
            f"""
            UPDATE import_runs
            SET run_status = %s,
                imported_rows = %s,
                skipped_rows = %s,
                error_rows = %s,
                failure_message = %s,
                completed_at = {completed_at_sql},
                last_activity_at = CURRENT_TIMESTAMP,
                updated_at = CURRENT_TIMESTAMP
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            """,
            (
                next_status,
                imported_rows,
                skipped_rows,
                error_rows,
                failure_message,
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        conn.commit()

        return {
            "import_run_id": import_run_id,
            "run_status": next_status,
            "imported_rows": imported_rows,
            "skipped_rows": skipped_rows,
            "error_rows": error_rows,
            "processed_this_packet": processed_this_packet,
            "remaining_rows": remaining_rows,
        }

    except Exception:
        conn.rollback()
        raise

    finally:
        cur.close()
        conn.close()


# =========================================================
# IMPORT RUN MAPPING DATA
# =========================================================

def get_import_run_mapping_data(
    import_run_id,
    *,
    spa_id,
    business_unit_id,
    preview_limit=10,
):
    """
    Load one durable Import Run for the column-mapping screen.

    Returns:
      - run metadata
      - source headers
      - PSP destination fields
      - current/suggested mapping
      - a small source-data preview

    This is read-only and remains strictly workspace scoped.
    """

    try:
        preview_limit = int(preview_limit)
    except (TypeError, ValueError):
        preview_limit = 10

    preview_limit = max(
        1,
        min(
            preview_limit,
            25,
        ),
    )

    conn = get_db_connection()
    cur = conn.cursor()

    try:
        cur.execute(
            """
            SELECT
                entity_type,
                run_status,
                source_filename,
                source_extension,
                source_size_bytes,
                mapping_json,
                options_json,
                total_rows,
                valid_rows,
                invalid_rows,
                strong_duplicate_rows,
                possible_duplicate_rows,
                imported_rows,
                skipped_rows,
                error_rows,
                requested_at
            FROM import_runs
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            LIMIT 1
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
            ),
        )

        run = cur.fetchone()

        if not run:
            raise ImportServiceError(
                "Import Run was not found in this workspace."
            )

        entity_type = str(
            run[0] or ""
        ).strip().lower()

        options_payload = (
            run[6]
            if isinstance(run[6], dict)
            else {}
        )

        profile = get_workspace_import_profile(
            cur,
            spa_id=spa_id,
            business_unit_id=business_unit_id,
            entity_type=entity_type,
            options=options_payload,
        )

        mapping_json = (
            run[5]
            if isinstance(run[5], dict)
            else {}
        )

        headers = mapping_json.get(
            "headers",
            [],
        )

        mapping = (
            mapping_json.get("selected")
            or mapping_json.get("suggestions")
            or []
        )

        cur.execute(
            """
            SELECT
                source_row_number,
                source_data
            FROM import_run_rows
            WHERE import_run_id = %s
              AND spa_id = %s
              AND business_unit_id = %s
            ORDER BY source_row_number
            LIMIT %s
            """,
            (
                import_run_id,
                spa_id,
                business_unit_id,
                preview_limit,
            ),
        )

        preview_rows = []

        for (
            source_row_number,
            source_data,
        ) in cur.fetchall():
            source_data = (
                source_data
                if isinstance(source_data, dict)
                else {}
            )

            preview_rows.append({
                "row_number": source_row_number,
                "values": source_data.get(
                    "values",
                    [],
                ),
            })

        return {
            "import_run_id": import_run_id,
            "entity_type": entity_type,
            "display_name": profile["display_name"],
            "run_status": run[1],
            "filename": run[2],
            "extension": run[3],
            "source_size_bytes": run[4],
            "headers": headers,
            "mapping": mapping,
            "options": options_payload,
            "fields": [
                {
                    "key": field["key"],
                    "label": field["label"],
                    "required": bool(
                        field.get("required")
                    ),
                }
                for field in profile["fields"]
            ],
            "preview_rows": preview_rows,
            "total_rows": run[7],
            "valid_rows": run[8],
            "invalid_rows": run[9],
            "strong_duplicate_rows": run[10],
            "possible_duplicate_rows": run[11],
            "imported_rows": run[12],
            "skipped_rows": run[13],
            "error_rows": run[14],
            "requested_at": run[15],
        }

    finally:
        cur.close()
        conn.close()
